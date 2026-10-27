"""Build the sample data dictionaries in examples/sample-dictionaries/.

Writes one CSV per source system plus a combined Excel workbook (one sheet per
source), so both import paths can be tested against the same metadata.

    uv run --no-project --with openpyxl python tools/build_sample_dictionaries.py
"""

from __future__ import annotations

import csv
from pathlib import Path

OUT = Path(__file__).resolve().parent.parent / "examples" / "sample-dictionaries"

FIELDS = [
    "source_system",
    "table_name",
    "table_description",
    "update_pattern",
    "column_name",
    "ordinal",
    "logical_type",
    "source_type",
    "nullable",
    "primary_key",
    "foreign_key",
    "accepted_values",
    "unit",
    "pii",
    "classification",
    "description",
    "example",
    "business_term",
    "owner",
]


def col(name, logical, source, desc, *, nullable=True, pk=False, fk="", accepted="", unit="",
        pii=False, classification="internal", example="", term=""):
    return dict(column_name=name, logical_type=logical, source_type=source, description=desc,
                nullable="Y" if nullable else "N", primary_key="Y" if pk else "N", foreign_key=fk,
                accepted_values=accepted, unit=unit, pii="Y" if pii else "N",
                classification=classification, example=example, business_term=term)


def table(source, name, desc, pattern, owner, columns):
    rows = []
    for i, c in enumerate(columns, start=1):
        rows.append(dict(source_system=source, table_name=name, table_description=desc,
                         update_pattern=pattern, ordinal=i, owner=owner, **c))
    return rows


# NYC TLC yellow taxi trips: monthly Parquet files plus the taxi zone lookup CSV.
# https://www.nyc.gov/site/tlc/about/tlc-trip-record-data.page
TLC = [
    *table("nyc_tlc", "yellow_tripdata",
           "One row per completed yellow taxi trip, published monthly as Parquet.",
           "monthly_file_append", "mobility-data", [
        col("VendorID", "integer", "INTEGER", "TPEP provider that supplied the record.",
            accepted="1=Creative Mobile Technologies;2=Curb Mobility;6=Myle Technologies;7=Helix",
            example="2", term="Taxi technology vendor"),
        col("tpep_pickup_datetime", "timestamp", "TIMESTAMP", "Date and time the meter was engaged (local New York time).",
            nullable=False, example="2024-01-01 00:57:55", term="Pickup time"),
        col("tpep_dropoff_datetime", "timestamp", "TIMESTAMP", "Date and time the meter was disengaged (local New York time).",
            nullable=False, example="2024-01-01 01:17:43", term="Dropoff time"),
        col("passenger_count", "integer", "BIGINT", "Number of passengers, entered by the driver.",
            example="1", term="Passenger count"),
        col("trip_distance", "decimal(10,2)", "DOUBLE", "Elapsed trip distance reported by the taximeter.",
            unit="miles", example="1.72", term="Trip distance"),
        col("RatecodeID", "integer", "BIGINT", "Final rate code in effect at the end of the trip.",
            accepted="1=Standard rate;2=JFK;3=Newark;4=Nassau or Westchester;5=Negotiated fare;6=Group ride;99=Unknown",
            example="1", term="Rate code"),
        col("store_and_fwd_flag", "string", "VARCHAR",
            "Y if the record was held in vehicle memory before sending because the vehicle had no server connection.",
            accepted="Y;N", example="N"),
        col("PULocationID", "integer", "INTEGER", "TLC taxi zone where the meter was engaged.",
            fk="taxi_zone_lookup.LocationID", example="186", term="Pickup zone"),
        col("DOLocationID", "integer", "INTEGER", "TLC taxi zone where the meter was disengaged.",
            fk="taxi_zone_lookup.LocationID", example="79", term="Dropoff zone"),
        col("payment_type", "integer", "BIGINT", "How the passenger paid for the trip.",
            accepted="0=Flex fare;1=Credit card;2=Cash;3=No charge;4=Dispute;5=Unknown;6=Voided trip",
            example="1", term="Payment method"),
        col("fare_amount", "decimal(10,2)", "DOUBLE", "Time-and-distance fare calculated by the meter.",
            unit="USD", classification="confidential", example="17.70", term="Base fare"),
        col("extra", "decimal(10,2)", "DOUBLE", "Miscellaneous extras and surcharges, such as rush-hour and overnight charges.",
            unit="USD", classification="confidential", example="1.00"),
        col("mta_tax", "decimal(10,2)", "DOUBLE", "MTA tax triggered automatically by the metered rate in use.",
            unit="USD", classification="confidential", example="0.50"),
        col("tip_amount", "decimal(10,2)", "DOUBLE", "Tip amount; populated for credit card tips only, cash tips are not recorded.",
            unit="USD", classification="confidential", example="3.72", term="Tip"),
        col("tolls_amount", "decimal(10,2)", "DOUBLE", "Total tolls paid during the trip.",
            unit="USD", classification="confidential", example="0.00"),
        col("improvement_surcharge", "decimal(10,2)", "DOUBLE", "Improvement surcharge assessed at flag drop.",
            unit="USD", classification="confidential", example="1.00"),
        col("total_amount", "decimal(10,2)", "DOUBLE", "Total amount charged to passengers; excludes cash tips.",
            unit="USD", classification="confidential", example="22.70", term="Trip revenue"),
        col("congestion_surcharge", "decimal(10,2)", "DOUBLE", "New York State congestion surcharge collected for the trip.",
            unit="USD", classification="confidential", example="2.50"),
        col("Airport_fee", "decimal(10,2)", "DOUBLE", "Fee for pickups at LaGuardia and JFK airports.",
            unit="USD", classification="confidential", example="0.00"),
    ]),
    *table("nyc_tlc", "taxi_zone_lookup",
           "Reference list of TLC taxi zones, used to decode pickup and dropoff location IDs.",
           "full_refresh_rarely_changes", "mobility-data", [
        col("LocationID", "integer", "BIGINT", "Taxi zone identifier.", nullable=False, pk=True,
            example="186", term="Taxi zone"),
        col("Borough", "string", "VARCHAR", "New York City borough containing the zone.",
            accepted="Manhattan;Brooklyn;Queens;Bronx;Staten Island;EWR;Unknown;N/A", example="Manhattan", term="Borough"),
        col("Zone", "string", "VARCHAR", "Neighbourhood or area name of the zone.", example="Penn Station/Madison Sq West"),
        col("service_zone", "string", "VARCHAR", "TLC service area the zone belongs to.",
            accepted="Yellow Zone;Boro Zone;Airports;EWR;N/A", example="Yellow Zone"),
    ]),
]

# Open-Meteo historical weather API (no key needed), flattened to one row per location and hour/day.
# https://open-meteo.com/en/docs/historical-weather-api
WEATHER_COMMON = [
    col("location_key", "string", "request parameter",
        "Our identifier for the requested point, e.g. a borough centroid; the API itself is queried by coordinates.",
        nullable=False, pk=True, example="manhattan", term="Weather location"),
    col("latitude", "decimal(9,6)", "number", "Latitude of the weather grid cell the API resolved to.",
        nullable=False, unit="degrees", example="40.738136"),
    col("longitude", "decimal(9,6)", "number", "Longitude of the weather grid cell the API resolved to.",
        nullable=False, unit="degrees", example="-74.04254"),
    col("timezone", "string", "string", "IANA time zone used for the time values.", nullable=False,
        example="America/New_York"),
]
WEATHER = [
    *table("open_meteo", "weather_hourly",
           "Hourly historical weather observations (reanalysis) per location from the Open-Meteo archive API.",
           "api_daily_incremental", "weather-data", [
        *WEATHER_COMMON,
        col("observed_at", "timestamp", "string (iso8601)", "Start of the hour, in local time for the chosen time zone.",
            nullable=False, pk=True, example="2024-01-01T00:00", term="Observation hour"),
        col("temperature_2m", "decimal(5,1)", "number", "Air temperature 2 metres above ground.", unit="°C", example="1.9",
            term="Temperature"),
        col("relative_humidity_2m", "integer", "number", "Relative humidity 2 metres above ground.", unit="%", example="74"),
        col("precipitation", "decimal(6,1)", "number", "Total precipitation (rain, showers, snow) in the preceding hour.",
            unit="mm", example="0.0", term="Precipitation"),
        col("rain", "decimal(6,1)", "number", "Rain from large-scale weather systems in the preceding hour.", unit="mm",
            example="0.0"),
        col("snowfall", "decimal(6,2)", "number", "Snowfall in the preceding hour.", unit="cm", example="0.00"),
        col("snow_depth", "decimal(6,2)", "number", "Snow depth on the ground.", unit="m", example="0.00"),
        col("weather_code", "integer", "number", "Weather condition as a WMO weather interpretation code.",
            accepted="0=Clear sky;1=Mainly clear;2=Partly cloudy;3=Overcast;45=Fog;51=Light drizzle;61=Slight rain;63=Moderate rain;65=Heavy rain;71=Slight snow;73=Moderate snow;75=Heavy snow;95=Thunderstorm",
            unit="wmo code", example="3", term="Weather condition"),
        col("cloud_cover", "integer", "number", "Total cloud cover as an area fraction.", unit="%", example="100"),
        col("wind_speed_10m", "decimal(5,1)", "number", "Wind speed 10 metres above ground.", unit="km/h", example="6.3"),
        col("wind_gusts_10m", "decimal(5,1)", "number", "Maximum wind gust 10 metres above ground in the preceding hour.",
            unit="km/h", example="13.0"),
    ]),
    *table("open_meteo", "weather_daily",
           "Daily aggregated historical weather per location from the Open-Meteo archive API.",
           "api_daily_incremental", "weather-data", [
        *WEATHER_COMMON,
        col("observed_date", "date", "string (iso8601)", "Calendar day in the chosen time zone.", nullable=False, pk=True,
            example="2024-01-01", term="Observation date"),
        col("temperature_2m_max", "decimal(5,1)", "number", "Maximum daily air temperature 2 metres above ground.",
            unit="°C", example="4.8"),
        col("temperature_2m_min", "decimal(5,1)", "number", "Minimum daily air temperature 2 metres above ground.",
            unit="°C", example="0.3"),
        col("precipitation_sum", "decimal(6,1)", "number", "Total precipitation for the day.", unit="mm", example="0.0"),
        col("snowfall_sum", "decimal(6,2)", "number", "Total snowfall for the day.", unit="cm", example="0.00"),
        col("weather_code", "integer", "number", "Most severe WMO weather code during the day.", unit="wmo code",
            example="3"),
    ]),
]

# TPC-H: a classic relational order-management schema, generated locally with DuckDB (CALL dbgen(sf=1)).
# https://www.tpc.org/tpch/  Names, addresses and phone numbers are synthetic but flagged as PII
# so masking and classification features can be exercised.
ERP_OWNER = "commerce-data"
ERP = [
    *table("tpch_erp", "region", "The five world regions.", "full_refresh_rarely_changes", ERP_OWNER, [
        col("r_regionkey", "integer", "INTEGER", "Region identifier.", nullable=False, pk=True, example="1"),
        col("r_name", "string", "VARCHAR", "Region name.", nullable=False,
            accepted="AFRICA;AMERICA;ASIA;EUROPE;MIDDLE EAST", example="AMERICA", term="Region"),
        col("r_comment", "string", "VARCHAR", "Free-text comment.", example="hs use ironic, even requests."),
    ]),
    *table("tpch_erp", "nation", "The 25 nations customers and suppliers belong to.", "full_refresh_rarely_changes",
           ERP_OWNER, [
        col("n_nationkey", "integer", "INTEGER", "Nation identifier.", nullable=False, pk=True, example="24"),
        col("n_name", "string", "VARCHAR", "Nation name.", nullable=False, example="UNITED STATES", term="Country"),
        col("n_regionkey", "integer", "INTEGER", "Region the nation belongs to.", nullable=False,
            fk="region.r_regionkey", example="1"),
        col("n_comment", "string", "VARCHAR", "Free-text comment.", example="y final packages."),
    ]),
    *table("tpch_erp", "customer", "Customers who place orders.", "cdc", ERP_OWNER, [
        col("c_custkey", "bigint", "BIGINT", "Customer identifier.", nullable=False, pk=True, example="1",
            term="Customer"),
        col("c_name", "string", "VARCHAR", "Customer name.", nullable=False, pii=True, classification="restricted",
            example="Customer#000000001", term="Customer name"),
        col("c_address", "string", "VARCHAR", "Customer street address.", pii=True, classification="restricted",
            example="IVhzIApeRb ot,c,E"),
        col("c_nationkey", "integer", "INTEGER", "Nation the customer is based in.", nullable=False,
            fk="nation.n_nationkey", example="15"),
        col("c_phone", "string", "VARCHAR", "Customer phone number.", pii=True, classification="restricted",
            example="25-989-741-2988"),
        col("c_acctbal", "decimal(15,2)", "DECIMAL(15,2)", "Current account balance.", unit="USD",
            classification="confidential", example="711.56", term="Account balance"),
        col("c_mktsegment", "string", "VARCHAR", "Market segment the customer belongs to.",
            accepted="AUTOMOBILE;BUILDING;FURNITURE;HOUSEHOLD;MACHINERY", example="BUILDING", term="Market segment"),
        col("c_comment", "string", "VARCHAR", "Free-text comment.", example="to the even, regular platelets."),
    ]),
    *table("tpch_erp", "supplier", "Suppliers that provide parts.", "cdc", ERP_OWNER, [
        col("s_suppkey", "bigint", "BIGINT", "Supplier identifier.", nullable=False, pk=True, example="1",
            term="Supplier"),
        col("s_name", "string", "VARCHAR", "Supplier name.", nullable=False, example="Supplier#000000001"),
        col("s_address", "string", "VARCHAR", "Supplier street address.", classification="confidential",
            example="N kD4on9OM Ipw3,gf0JBoQDd7tgrzrddZ"),
        col("s_nationkey", "integer", "INTEGER", "Nation the supplier is based in.", nullable=False,
            fk="nation.n_nationkey", example="17"),
        col("s_phone", "string", "VARCHAR", "Supplier phone number.", classification="confidential",
            example="27-918-335-1736"),
        col("s_acctbal", "decimal(15,2)", "DECIMAL(15,2)", "Current account balance.", unit="USD",
            classification="confidential", example="5755.94"),
        col("s_comment", "string", "VARCHAR", "Free-text comment.", example="each slyly above the careful"),
    ]),
    *table("tpch_erp", "part", "Parts sold to customers.", "cdc", ERP_OWNER, [
        col("p_partkey", "bigint", "BIGINT", "Part identifier.", nullable=False, pk=True, example="1", term="Product"),
        col("p_name", "string", "VARCHAR", "Part name.", nullable=False, example="goldenrod lavender spring chocolate lace"),
        col("p_mfgr", "string", "VARCHAR", "Manufacturer.", example="Manufacturer#1"),
        col("p_brand", "string", "VARCHAR", "Brand.", example="Brand#13", term="Brand"),
        col("p_type", "string", "VARCHAR", "Part type.", example="PROMO BURNISHED COPPER"),
        col("p_size", "integer", "INTEGER", "Part size.", example="7"),
        col("p_container", "string", "VARCHAR", "Shipping container type.", example="JUMBO PKG"),
        col("p_retailprice", "decimal(15,2)", "DECIMAL(15,2)", "List price.", unit="USD", example="901.00",
            term="List price"),
        col("p_comment", "string", "VARCHAR", "Free-text comment.", example="ly. slyly ironi"),
    ]),
    *table("tpch_erp", "partsupp", "Which supplier offers which part, with stock and cost.", "cdc", ERP_OWNER, [
        col("ps_partkey", "bigint", "BIGINT", "Part supplied.", nullable=False, pk=True, fk="part.p_partkey",
            example="1"),
        col("ps_suppkey", "bigint", "BIGINT", "Supplier offering the part.", nullable=False, pk=True,
            fk="supplier.s_suppkey", example="2"),
        col("ps_availqty", "integer", "BIGINT", "Quantity available from this supplier.", unit="units", example="3325"),
        col("ps_supplycost", "decimal(15,2)", "DECIMAL(15,2)", "Cost of the part from this supplier.", unit="USD",
            classification="confidential", example="771.64", term="Supply cost"),
        col("ps_comment", "string", "VARCHAR", "Free-text comment.", example="requests after the carefully ironic"),
    ]),
    *table("tpch_erp", "orders", "Customer orders (order header).", "cdc", ERP_OWNER, [
        col("o_orderkey", "bigint", "BIGINT", "Order identifier.", nullable=False, pk=True, example="1", term="Order"),
        col("o_custkey", "bigint", "BIGINT", "Customer who placed the order.", nullable=False,
            fk="customer.c_custkey", example="370"),
        col("o_orderstatus", "string", "VARCHAR", "Order status derived from its line items.",
            accepted="F=Fulfilled;O=Open;P=Partially fulfilled", example="O", term="Order status"),
        col("o_totalprice", "decimal(15,2)", "DECIMAL(15,2)", "Total order value including tax and discounts.",
            unit="USD", classification="confidential", example="172799.49", term="Order value"),
        col("o_orderdate", "date", "DATE", "Date the order was placed.", nullable=False, example="1996-01-02",
            term="Order date"),
        col("o_orderpriority", "string", "VARCHAR", "Order priority.",
            accepted="1-URGENT;2-HIGH;3-MEDIUM;4-NOT SPECIFIED;5-LOW", example="5-LOW"),
        col("o_clerk", "string", "VARCHAR", "Clerk who handled the order.", pii=True, classification="restricted",
            example="Clerk#000000951"),
        col("o_shippriority", "integer", "INTEGER", "Shipping priority.", example="0"),
        col("o_comment", "string", "VARCHAR", "Free-text comment.", example="nstructions sleep furiously among"),
    ]),
    *table("tpch_erp", "lineitem", "Order lines: one row per part on an order.", "cdc", ERP_OWNER, [
        col("l_orderkey", "bigint", "BIGINT", "Order the line belongs to.", nullable=False, pk=True,
            fk="orders.o_orderkey", example="1"),
        col("l_partkey", "bigint", "BIGINT", "Part ordered.", nullable=False, fk="part.p_partkey;partsupp.ps_partkey",
            example="1552"),
        col("l_suppkey", "bigint", "BIGINT", "Supplier fulfilling the line.", nullable=False,
            fk="supplier.s_suppkey;partsupp.ps_suppkey", example="93"),
        col("l_linenumber", "integer", "BIGINT", "Line number within the order.", nullable=False, pk=True, example="1"),
        col("l_quantity", "decimal(15,2)", "DECIMAL(15,2)", "Quantity ordered.", unit="units", example="17.00",
            term="Quantity"),
        col("l_extendedprice", "decimal(15,2)", "DECIMAL(15,2)", "Quantity times part price, before discount and tax.",
            unit="USD", classification="confidential", example="24710.35", term="Gross line amount"),
        col("l_discount", "decimal(15,2)", "DECIMAL(15,2)", "Discount rate applied to the line.", unit="ratio",
            example="0.04"),
        col("l_tax", "decimal(15,2)", "DECIMAL(15,2)", "Tax rate applied to the line.", unit="ratio", example="0.02"),
        col("l_returnflag", "string", "VARCHAR", "Whether the line was returned.", accepted="A;N;R", example="N"),
        col("l_linestatus", "string", "VARCHAR", "Line status.", accepted="F=Fulfilled;O=Open", example="O"),
        col("l_shipdate", "date", "DATE", "Date the line shipped.", example="1996-03-13", term="Ship date"),
        col("l_commitdate", "date", "DATE", "Date the line was promised to ship.", example="1996-02-12"),
        col("l_receiptdate", "date", "DATE", "Date the customer received the line.", example="1996-03-22"),
        col("l_shipinstruct", "string", "VARCHAR", "Shipping instructions.",
            accepted="DELIVER IN PERSON;COLLECT COD;NONE;TAKE BACK RETURN", example="DELIVER IN PERSON"),
        col("l_shipmode", "string", "VARCHAR", "Shipping mode.", accepted="AIR;FOB;MAIL;RAIL;REG AIR;SHIP;TRUCK",
            example="TRUCK", term="Ship mode"),
        col("l_comment", "string", "VARCHAR", "Free-text comment.", example="egular courts above the"),
    ]),
]

SOURCES = {"nyc_tlc": TLC, "open_meteo": WEATHER, "tpch_erp": ERP}


def main() -> None:
    OUT.mkdir(parents=True, exist_ok=True)
    for name, rows in SOURCES.items():
        with (OUT / f"{name}.csv").open("w", newline="", encoding="utf-8") as f:
            writer = csv.DictWriter(f, fieldnames=FIELDS)
            writer.writeheader()
            writer.writerows(rows)

    from openpyxl import Workbook
    from openpyxl.styles import Font

    wb = Workbook()
    wb.remove(wb.active)
    for name, rows in SOURCES.items():
        ws = wb.create_sheet(name)
        ws.append(FIELDS)
        for cell in ws[1]:
            cell.font = Font(bold=True)
        for row in rows:
            ws.append([row[f] for f in FIELDS])
        ws.freeze_panes = "A2"
    wb.save(OUT / "sample_dictionaries.xlsx")

    for name, rows in SOURCES.items():
        tables = {r["table_name"] for r in rows}
        print(f"{name}: {len(tables)} tables, {len(rows)} columns")


if __name__ == "__main__":
    main()
