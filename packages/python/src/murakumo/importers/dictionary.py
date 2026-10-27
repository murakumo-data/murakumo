"""Import a data dictionary: a CSV file or Excel workbook with one row per column.

The format is described in examples/README.md. Only ``table_name`` and ``column_name`` are
required; every other field is used when present. Headers are matched loosely, so ``Table``,
``table name`` and ``TABLE_NAME`` all work.
"""

from __future__ import annotations

import csv
import re
from collections import defaultdict
from pathlib import Path
from typing import Any

from murakumo.diagnostics import Diagnostic, error, warning
from murakumo.importers.base import SourceImport, fk_relationship, safe_name, unique_names
from murakumo.importers.typemap import is_logical_type, sql_to_logical
from murakumo.ir.model import Attribute, Classification, Entity, KeyRole, LoadPattern, Relationship, SourceSystem

ALIASES = {
    "source": "source_system",
    "system": "source_system",
    "table": "table_name",
    "entity": "table_name",
    "table_desc": "table_description",
    "column": "column_name",
    "field": "column_name",
    "field_name": "column_name",
    "attribute": "column_name",
    "type": "logical_type",
    "data_type": "source_type",
    "datatype": "source_type",
    "position": "ordinal",
    "null": "nullable",
    "is_nullable": "nullable",
    "pk": "primary_key",
    "is_primary_key": "primary_key",
    "fk": "foreign_key",
    "references": "foreign_key",
    "allowed_values": "accepted_values",
    "values": "accepted_values",
    "is_pii": "pii",
    "comment": "description",
    "definition": "description",
    "glossary_term": "business_term",
}

_TRUE = {"y", "yes", "true", "t", "1", "x"}
_FALSE = {"n", "no", "false", "f", "0", ""}


def _header(text: Any) -> str:
    key = re.sub(r"[^a-z0-9]+", "_", str(text or "").strip().lower()).strip("_")
    return ALIASES.get(key, key)


def read_rows(path: Path, sheet: str | None = None) -> list[tuple[str, int, dict[str, str]]]:
    """Rows as (sheet or file name, row number, normalised fields)."""
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        return _read_excel(path, sheet)
    with path.open(newline="", encoding="utf-8-sig") as handle:
        sample = handle.read(8192)
        handle.seek(0)
        try:
            dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
        except csv.Error:
            dialect = csv.excel
        reader = csv.reader(handle, dialect)
        header = [_header(h) for h in next(reader, [])]
        return [
            (path.name, number, {h: (v or "").strip() for h, v in zip(header, row, strict=False) if h})
            for number, row in enumerate(reader, start=2)
            if any(cell.strip() for cell in row)
        ]


def _read_excel(path: Path, sheet: str | None) -> list[tuple[str, int, dict[str, str]]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    if sheet is not None and sheet not in workbook.sheetnames:
        raise ValueError(f"no sheet {sheet!r} in {path.name}; sheets: {', '.join(workbook.sheetnames)}")
    rows: list[tuple[str, int, dict[str, str]]] = []
    for name in [sheet] if sheet else workbook.sheetnames:
        header: list[str] | None = None
        for number, values in enumerate(workbook[name].iter_rows(values_only=True), start=1):
            cells = ["" if v is None else str(v).strip() for v in values]
            if not any(cells):
                continue
            if header is None:
                header = [_header(c) for c in cells]
                if "table_name" not in header or "column_name" not in header:
                    header = None  # not a dictionary sheet (a cover page, notes...)
                    break
                continue
            rows.append((name, number, {h: v for h, v in zip(header, cells, strict=False) if h}))
    workbook.close()
    return rows


def _bool(value: str) -> bool | None:
    lowered = value.strip().lower()
    if lowered in _TRUE:
        return True
    if lowered in _FALSE:
        return False
    return None


def _load_pattern(pattern: str) -> LoadPattern | None:
    lowered = pattern.lower()
    for needle, load in (
        ("cdc", LoadPattern.CDC),
        ("stream", LoadPattern.STREAMING),
        ("append", LoadPattern.APPEND),
        ("increment", LoadPattern.INCREMENTAL),
        ("full", LoadPattern.FULL),
        ("refresh", LoadPattern.FULL),
        ("snapshot", LoadPattern.FULL),
    ):
        if needle in lowered:
            return load
    return None


def _accepted_values(text: str) -> dict[str, str | None]:
    values: dict[str, str | None] = {}
    for part in re.split(r"[;|]", text):
        part = part.strip()
        if not part:
            continue
        code, sep, label = part.partition("=")
        values[code.strip()] = label.strip() if sep else None
    return values


def _foreign_keys(text: str) -> list[tuple[str, str]]:
    """``orders.o_orderkey; part(p_partkey)`` -> [("orders", "o_orderkey"), ("part", "p_partkey")]."""
    refs = []
    for part in re.split(r"[;,|]", text):
        part = part.strip()
        match = re.fullmatch(r"([^.()\s]+(?:\.[^.()\s]+)?)\s*(?:\.([^.()\s]+)|\(\s*([^()\s]+)\s*\))", part)
        if match:
            table = match.group(1).split(".")[-1]  # schema.table.column -> table
            refs.append((table, match.group(2) or match.group(3)))
        elif part:
            refs.append(("", part))
    return refs


def import_dictionary(path: Path, source: str | None = None, sheet: str | None = None) -> list[SourceImport]:
    """Read a dictionary into one SourceImport per source system it mentions."""
    rows = read_rows(path, sheet)
    by_source: dict[str, dict[str, list[tuple[str, int, dict[str, str]]]]] = defaultdict(lambda: defaultdict(list))
    diagnostics: list[Diagnostic] = []
    if not rows:
        diagnostics.append(error("MK402", "no rows found; is this a data dictionary?", path))

    for origin, number, row in rows:
        where = f"{origin} row {number}" if origin != path.name else f"row {number}"
        if not row.get("table_name") or not row.get("column_name"):
            diagnostics.append(error("MK402", "row has no table_name or column_name; skipped", path, where))
            continue
        system = safe_name(source or row.get("source_system") or path.stem, "source")
        by_source[system][row["table_name"]].append((where, number, row))

    imports: list[SourceImport] = []
    for system, tables in by_source.items():
        result = SourceImport(source=SourceSystem(name=system))
        foreign: list[tuple[str, str, str, str, bool, str]] = []  # from_table, column, to_table, to_column, nullable
        pks: dict[str, set[str]] = {}

        for table_name, table_rows in tables.items():
            entity_name = safe_name(table_name, "table")
            first = table_rows[0][2]
            fields: dict[str, Any] = {"name": entity_name, "source": system}
            if first.get("table_description"):
                fields["description"] = first["table_description"]
            if first.get("owner"):
                fields["owner"] = first["owner"]
            if first.get("update_pattern"):
                load = _load_pattern(first["update_pattern"])
                if load:
                    fields["load"] = load
                fields["properties"] = {"update_pattern": first["update_pattern"]}

            ordered = sorted(table_rows, key=lambda r: (_ordinal(r[2].get("ordinal")), r[1]))
            attributes = []
            for where, _number, row in ordered:
                column, problems = _attribute(row)
                for level, message in problems:
                    diagnostics.append((error if level == "error" else warning)("MK403", message, path, where))
                if column is None:
                    continue
                attributes.append(column)
                for to_table, to_column in _foreign_keys(row.get("foreign_key", "")):
                    if not to_table:
                        diagnostics.append(
                            warning("MK403", f"cannot read foreign key {to_column!r}; use table.column", path, where)
                        )
                        continue
                    foreign.append(
                        (entity_name, column.name, safe_name(to_table), safe_name(to_column), column.nullable, where)
                    )
            fields["attributes"] = attributes
            entity = Entity(**fields)
            pks[entity_name.lower()] = {a.name.lower() for a in entity.primary_key}
            result.entities.append(entity)

        tables_here = {e.name.lower() for e in result.entities}
        for _f, _c, to_table, _tc, _n, where in foreign:
            if to_table.lower() not in tables_here:
                diagnostics.append(
                    warning("MK403", f"foreign key points to {to_table}, which is not in this dictionary", path, where)
                )
        result.relationships = _relationships(foreign, pks)
        imports.append(result)

    if imports:
        imports[0].diagnostics.extend(diagnostics)
    elif diagnostics:
        imports.append(SourceImport(source=SourceSystem(name=safe_name(source or path.stem)), diagnostics=diagnostics))
    return imports


def _ordinal(value: str | None) -> float:
    try:
        return float(value) if value else float("inf")
    except ValueError:
        return float("inf")


def _attribute(row: dict[str, str]) -> tuple[Attribute | None, list[tuple[str, str]]]:
    problems: list[tuple[str, str]] = []
    name = safe_name(row["column_name"])
    if name != row["column_name"]:
        problems.append(("warning", f"column {row['column_name']!r} renamed to {name!r} to make a valid name"))

    logical = row.get("logical_type", "")
    source_type = row.get("source_type", "")
    if logical and is_logical_type(logical):
        type_ = logical
    else:
        mapped, exact = sql_to_logical(logical or source_type or "string")
        type_ = mapped
        if not exact:
            problems.append(("warning", f"unknown type {logical or source_type!r} for {name}; using string"))

    fields: dict[str, Any] = {"name": name, "type": type_}
    if source_type:
        fields["source_type"] = source_type
    primary = _bool(row.get("primary_key", ""))
    if primary:
        fields["key"] = KeyRole.PRIMARY
    nullable = _bool(row.get("nullable", ""))
    if nullable is not None and row.get("nullable", "").strip():
        fields["nullable"] = nullable
    elif primary:
        fields["nullable"] = False
    if row.get("accepted_values"):
        fields["accepted_values"] = _accepted_values(row["accepted_values"])
    if _bool(row.get("pii", "")):
        fields["pii"] = True
    if row.get("classification"):
        try:
            fields["classification"] = Classification(row["classification"].strip().lower())
        except ValueError:
            problems.append(("warning", f"unknown classification {row['classification']!r} for {name}; ignored"))
    for key in ("description", "unit", "example", "business_term"):
        if row.get(key):
            fields[key] = row[key]
    try:
        return Attribute(**fields), problems
    except ValueError as exc:
        problems.append(("error", f"column {name}: {exc}"))
        return None, problems


def _relationships(foreign: list[tuple], pks: dict[str, set[str]]) -> list[Relationship]:
    """Group foreign key columns into relationships.

    Columns pointing at the same table form one composite relationship only when together
    they are that table's whole primary key (lineitem -> partsupp on part and supplier).
    Otherwise each column is its own relationship (a trip's pickup zone and drop-off zone).
    """
    grouped: dict[tuple[str, str], list[tuple[str, str, bool]]] = defaultdict(list)
    for from_table, column, to_table, to_column, nullable, _where in foreign:
        grouped[(from_table, to_table)].append((column, to_column, nullable))

    relationships = []
    for (from_table, to_table), pairs in grouped.items():
        target_pk = pks.get(to_table.lower(), set())
        to_columns = [p[1].lower() for p in pairs]
        if len(pairs) > 1 and len(target_pk) > 1 and sorted(to_columns) == sorted(target_pk):
            relationships.append(
                fk_relationship(from_table, to_table, [(a, b) for a, b, _ in pairs], any(n for _, _, n in pairs))
            )
        else:
            for column, to_column, nullable in pairs:
                relationships.append(fk_relationship(from_table, to_table, [(column, to_column)], nullable))
    return unique_names(relationships)
